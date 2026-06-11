"""文档文本提取工具：把附件尽量转成模型可理解的纯文本。

这是多模态输入链路里很关键的一层：
- 图片通常保留给视觉模型处理
- PDF / DOCX / XLSX / PPTX / 纯文本文件则尽量抽出正文
- 抽出的文本会被拼接回用户输入上下文，供 LLM 阅读
"""

import mimetypes
from pathlib import Path

from loguru import logger

from nanobot.utils.helpers import detect_image_mime

# 当前支持尝试提取文本的扩展名集合。
SUPPORTED_EXTENSIONS: set[str] = {
    # 文档格式
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    # 纯文本/结构化文本格式
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".log",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    # 图片格式（当前主要作为占位；未来可扩展 OCR）
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
}

_MAX_TEXT_LENGTH = 200_000


def extract_text(path: Path) -> str | None:
    """根据文件类型选择合适的提取策略。

    返回值约定：
    - 成功：提取出的文本
    - 不支持的类型：``None``
    - 出错：形如 ``[error: ...]`` 的错误文本

    这样调用方可以不抛异常地继续拼装上下文。
    """
    if not isinstance(path, Path):
        path = Path(path)

    if not path.exists():
        return f"[error: file not found: {path}]"

    ext = path.suffix.lower()

    # 各格式解析器都采用延迟导入，避免进程启动时一次性加载大量文档库。
    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext == ".xlsx":
        return _extract_xlsx(path)
    elif ext == ".pptx":
        return _extract_pptx(path)
    elif _is_text_extension(ext):
        return _extract_text_file(path)
    elif ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        # 图片暂时不在这里做 OCR，只返回一个可读占位。
        return f"[image: {path.name}]"
    else:
        # 当前未支持的扩展名直接返回 None，让上层决定如何处理。
        return None


def _extract_pdf(path: Path) -> str:
    """使用 ``pypdf`` 提取 PDF 文本。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[error: pypdf not installed]"
    try:
        reader = PdfReader(path)
        pages: list[str] = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"--- Page {i} ---\n{text}")
        return _truncate("\n\n".join(pages), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PDF {}", path)
        return f"[error: failed to extract PDF: {e!s}]"


def _extract_docx(path: Path) -> str:
    """使用 ``python-docx`` 提取 DOCX 文本。"""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        return "[error: python-docx not installed]"
    try:
        doc = DocxDocument(path)
        paragraphs: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        return _truncate("\n\n".join(paragraphs), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract DOCX {}", path)
        return f"[error: failed to extract DOCX: {e!s}]"


def _extract_xlsx(path: Path) -> str:
    """使用 ``openpyxl`` 提取 XLSX 中的单元格文本。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[error: openpyxl not installed]"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[str] = []
                for row in ws.iter_rows(values_only=True):
                    row_text = "\t".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        rows.append(row_text)
                if rows:
                    sheets.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
            return _truncate("\n\n".join(sheets), _MAX_TEXT_LENGTH)
        finally:
            wb.close()
    except Exception as e:
        logger.exception("Failed to extract XLSX {}", path)
        return f"[error: failed to extract XLSX: {e!s}]"


def _extract_pptx(path: Path) -> str:
    """使用 ``python-pptx`` 提取 PPTX 中的幻灯片文本。"""
    try:
        from pptx import Presentation as PptxPresentation
    except ImportError:
        return "[error: python-pptx not installed]"
    try:
        prs = PptxPresentation(path)
        slides: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                _collect_pptx_shape_text(shape, slide_text)
            if slide_text:
                slides.append(f"--- Slide {i} ---\n" + "\n".join(slide_text))
        return _truncate("\n\n".join(slides), _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to extract PPTX {}", path)
        return f"[error: failed to extract PPTX: {e!s}]"


def _collect_pptx_shape_text(shape, out: list[str]) -> None:
    """递归收集一个 PPTX shape 里的文本。

    PPTX 里的“可见内容”不一定都在 ``shape.text``：
    - 组合图形需要继续遍历 ``.shapes``
    - 表格内容要从 ``.table`` 里读 cell 文本
    """
    sub_shapes = getattr(shape, "shapes", None)
    if sub_shapes is not None:
        for sub in sub_shapes:
            _collect_pptx_shape_text(sub, out)
        return

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                out.append(line)
        return

    text = getattr(shape, "text", "")
    if text:
        out.append(text)


def _extract_text_file(path: Path) -> str:
    """读取纯文本文件内容。"""
    try:
        # 先尝试 UTF-8；失败时退回 latin-1，尽量不要因为编码问题整份文件都读不了。
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1")
        return _truncate(content, _MAX_TEXT_LENGTH)
    except Exception as e:
        logger.exception("Failed to read text file {}", path)
        return f"[error: failed to read file: {e!s}]"


def _truncate(text: str, max_length: int) -> str:
    """按长度上限截断文本，并补一个说明后缀。"""
    if len(text) <= max_length:
        return text
    return text[:max_length] + f"... (truncated, {len(text)} chars total)"


def _is_text_extension(ext: str) -> bool:
    """判断扩展名是否属于直接按文本读取的格式。"""
    return ext in {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".xml",
        ".html",
        ".htm",
        ".log",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
    }


# ---------------------------------------------------------------------------
# 更高层的辅助函数：把附件拆成“交给视觉模型的图片”和“可直接抽文本的文档”。
# ---------------------------------------------------------------------------

_MAX_EXTRACT_FILE_SIZE = 50 * 1024 * 1024  # 50 MB


def is_image_file(path: str) -> bool:
    """判断一个文件是否像图片。

    优先通过文件头 magic bytes 判断；如果拿不到，再退回扩展名 / mimetype 猜测。
    """
    p = Path(path)
    mime: str | None = None
    if p.is_file():
        try:
            with p.open("rb") as f:
                mime = detect_image_mime(f.read(16))
        except OSError:
            mime = None
    if not mime:
        mime = mimetypes.guess_type(path)[0]
    return bool(mime and mime.startswith("image/"))


def reference_non_image_attachments(
    content: str, media: list[str],
) -> tuple[str, list[str]]:
    """把附件分成图片和非图片，但不读取非图片正文。

    适合只想保留“这个附件存在”的引用信息，而不做全文提取的场景。
    """
    image_paths: list[str] = []
    attachment_refs: list[str] = []
    for path in media:
        if is_image_file(path):
            image_paths.append(path)
        else:
            attachment_refs.append(f"[Attachment: {path}]")
    if attachment_refs:
        suffix = "\n".join(attachment_refs)
        content = f"{content}\n\n{suffix}" if content else suffix
    return content, image_paths


def extract_documents(
    text: str,
    media_paths: list[str],
    *,
    max_file_size: int = _MAX_EXTRACT_FILE_SIZE,
) -> tuple[str, list[str]]:
    """把附件拆成“图片”和“可提取文本的文档”。

    处理结果：
    - 图片路径继续保留，交给后续视觉块构造逻辑
    - 文档文本会被提取并拼接到 ``text`` 末尾

    之所以限制文件大小，是为了避免单个超大附件把内存或 CPU 吃爆。
    """
    image_paths: list[str] = []
    doc_texts: list[str] = []

    for path_str in media_paths:
        p = Path(path_str)
        if not p.is_file():
            continue

        try:
            size = p.stat().st_size
        except OSError:
            continue
        if size > max_file_size:
            logger.warning(
                "Skipping oversized file for extraction: {} ({:.1f} MB > {} MB limit)",
                p.name, size / (1024 * 1024), max_file_size // (1024 * 1024),
            )
            continue

        if is_image_file(path_str):
            image_paths.append(path_str)
        else:
            extracted = extract_text(p)
            if extracted and not extracted.startswith("[error:"):
                doc_texts.append(f"[File: {p.name}]\n{extracted}")

    if doc_texts:
        text = text + "\n\n" + "\n\n".join(doc_texts)

    return text, image_paths
